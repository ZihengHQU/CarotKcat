"""Build the submission-ready CarotKcat manuscript DOCX from the markdown source.

Parses a controlled markdown subset:
  # title / ## heading / ### subheading
  paragraphs with **bold** and *italic* inline marks
  ![alt](path) image lines (path relative to this script's directory)
  | a | b | pipe tables (rendered as three-line journal tables)
  - bullets, 1. numbered items
Journal style: Times New Roman 11, 1.5 line spacing, justified body,
three-line tables, centred figures with no borders, page-numbered footer.
"""
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(BASE, "CarotKcat_manuscript.md")
OUT = os.path.join(BASE, "CarotKcat_manuscript_v2.docx")
FONT = "Times New Roman"

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def add_runs(par, text):
    """Add runs to a paragraph, honouring **bold** and *italic*."""
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            r = par.add_run(chunk[2:-2]); r.bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            r = par.add_run(chunk[1:-1]); r.italic = True
        else:
            par.add_run(chunk)


def set_par_style(par, size=11, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                  space_after=6, line=1.5):
    par.alignment = align
    pf = par.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    for r in par.runs:
        r.font.name = FONT
        r.font.size = Pt(size)
        if bold:
            r.bold = True
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "bottom", "left", "right"):
        spec = kwargs.get(edge)
        tag = "w:" + edge
        el = tcBorders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            tcBorders.append(el)
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec))
            el.set(qn("w:color"), "000000")


def add_three_line_table(doc, rows):
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    last = len(rows) - 1
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            add_runs(par, cell_text)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else (
                WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            for r in par.runs:
                r.font.name = FONT
                r.font.size = Pt(9)
                if i == 0:
                    r.bold = True
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
            top = 12 if i == 0 else (6 if i == 1 else None)
            bottom = 12 if i == last else None
            set_cell_border(cell, top=top, bottom=bottom, left=None, right=None)
    return table


def add_page_number_footer(doc):
    footer = doc.sections[0].footer
    par = footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for tag, attr, val in (("w:fldChar", "w:fldCharType", "begin"),
                           ("w:instrText", "xml:space", "preserve"),
                           ("w:fldChar", "w:fldCharType", "end")):
        el = OxmlElement(tag)
        el.set(qn(attr), val)
        if tag == "w:instrText":
            el.text = "PAGE"
        run._r.append(el)
    run.font.name = FONT
    run.font.size = Pt(9)


def flush_table(doc, buf):
    rows = []
    for line in buf:
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        rows.append(cells)
    if rows:
        add_three_line_table(doc, rows)


def main():
    lines = open(SRC, encoding="utf-8").read().splitlines()
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)
    add_page_number_footer(doc)

    i = 0
    table_buf = []
    first_h1_done = False
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        if table_buf:
            flush_table(doc, table_buf)
            table_buf = []
        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).replace("**", "")
            if level == 1 and not first_h1_done:
                # manuscript title
                par = doc.add_paragraph()
                add_runs(par, text)
                set_par_style(par, size=16, bold=True,
                              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14, line=1.3)
                first_h1_done = True
            else:
                par = doc.add_paragraph()
                r = par.add_run(text)
                r.bold = True
                r.font.name = FONT
                r.font.size = Pt(13 if level <= 2 else 11.5)
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "SimHei")
                par.paragraph_format.space_before = Pt(12)
                par.paragraph_format.space_after = Pt(6)
            i += 1
            continue

        m = re.match(r"^!\[[^\]]*\]\(([^)]+)\)", line.strip())
        if m:
            path = os.path.normpath(os.path.join(BASE, m.group(1)))
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = par.add_run()
            width = Cm(15.5)
            if "fig_scatter" in path:
                width = Cm(12.5)
            if "fig_architecture" in path:
                width = Cm(14.5)
            run.add_picture(path, width=width)
            par.paragraph_format.space_after = Pt(10)
            i += 1
            continue

        stripped = line.strip()
        if stripped.startswith("- "):
            par = doc.add_paragraph(style="List Bullet")
            add_runs(par, stripped[2:])
            set_par_style(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=3)
            i += 1
            continue
        if re.match(r"^\d+\.\s", stripped):
            par = doc.add_paragraph(style="List Number")
            add_runs(par, re.sub(r"^\d+\.\s", "", stripped))
            set_par_style(par, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=3)
            i += 1
            continue

        par = doc.add_paragraph()
        add_runs(par, stripped)
        set_par_style(par)
        i += 1

    if table_buf:
        flush_table(doc, table_buf)

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
